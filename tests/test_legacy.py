"""Backwards-compatibility test for the legacy `python generate.py` command.

The command must keep working with no AI packages and no API key, writing the
same files in the same place. Output is compared with tests/golden/ by content
and structure, not bytes, because docx/PDF metadata changes on every run.

To refresh the golden files after an intentional template or sample change, run
the legacy command on tests/fixtures/data/sample.json and copy output/sample/*
into tests/golden/.
"""

import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parent.parent
GOLDEN = REPO / "tests" / "golden"
FIXTURE_DATA = REPO / "tests" / "fixtures" / "data"
EXPECTED_FILES = {
    "CV_AdrianMatei.docx",
    "CV_AdrianMatei.pdf",
    "CV_JohnDoe.docx",
    "CV_JohnDoe.pdf",
}
AI_MODULES = ("anthropic", "pydantic", "fitz", "pymupdf", "dotenv")


def _legacy_env():
    env = os.environ.copy()
    env.pop("ANTHROPIC_API_KEY", None)
    return env


def _workdir(tmp_path):
    """A working directory laid out like the repo: templates/ and data/ only."""
    shutil.copytree(REPO / "templates", tmp_path / "templates")
    shutil.copytree(FIXTURE_DATA, tmp_path / "data")
    return tmp_path


@pytest.fixture(scope="module")
def legacy_run(tmp_path_factory):
    work = _workdir(tmp_path_factory.mktemp("legacy"))
    result = subprocess.run(
        [sys.executable, str(REPO / "generate.py")],
        cwd=work,
        env=_legacy_env(),
        capture_output=True,
        text=True,
        timeout=300,
    )
    return work, result


# ---------------------------
# DOCX STRUCTURE
# ---------------------------
def _run_info(run):
    font = run.font
    color = font.color.rgb if font.color is not None and font.color.type is not None else None
    return (
        run.text,
        run.bold,
        run.italic,
        run.underline,
        font.size,
        font.name,
        str(color) if color is not None else None,
        run.style.name if run.style is not None else None,
    )


def _paragraph_info(p):
    content = []
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            content.append(("link", item.address, [_run_info(r) for r in item.runs]))
        else:
            content.append(("run",) + _run_info(item))
    return ("p", p.style.name, p.alignment, content)


def _block_info(element, parent):
    tag = element.tag.rsplit("}", 1)[-1]
    if tag == "p":
        return _paragraph_info(Paragraph(element, parent))
    if tag == "tbl":
        table = Table(element, parent)
        return (
            "tbl",
            [[[_block_info(child, cell) for child in cell._element.iterchildren()
               if child.tag.rsplit("}", 1)[-1] in ("p", "tbl")]
              for cell in row.cells]
             for row in table.rows],
        )
    return None


def docx_structure(path):
    doc = Document(str(path))
    body = [
        info for info in (_block_info(el, doc._body) for el in doc.element.body.iterchildren())
        if info is not None
    ]
    sections = [
        (
            s.page_width, s.page_height, s.orientation,
            s.left_margin, s.right_margin, s.top_margin, s.bottom_margin,
            [p.text for p in s.header.paragraphs],
            [p.text for p in s.footer.paragraphs],
        )
        for s in doc.sections
    ]
    links = sorted(
        rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")
    )
    return {"body": body, "sections": sections, "links": links}


def pdf_pages(path):
    import pymupdf  # test-only dependency; the legacy command never imports it

    with pymupdf.open(str(path)) as doc:
        return [(page.rect.width, page.rect.height, page.get_text()) for page in doc]


# ---------------------------
# TESTS
# ---------------------------
def test_legacy_command_succeeds(legacy_run):
    _, result = legacy_run
    assert result.returncode == 0, result.stdout + result.stderr
    assert "Processing: sample" in result.stdout
    assert "Done: sample" in result.stdout


def test_legacy_writes_same_files_in_same_place(legacy_run):
    work, _ = legacy_run
    written = {
        p.relative_to(work).as_posix()
        for p in (work / "output").rglob("*") if p.is_file()
    }
    assert written == {f"output/sample/{name}" for name in EXPECTED_FILES}


@pytest.mark.parametrize("name", ["CV_AdrianMatei.docx", "CV_JohnDoe.docx"])
def test_docx_matches_golden(legacy_run, name):
    work, _ = legacy_run
    actual = docx_structure(work / "output" / "sample" / name)
    expected = docx_structure(GOLDEN / name)
    assert actual["links"] == expected["links"]
    assert actual["sections"] == expected["sections"]
    assert len(actual["body"]) == len(expected["body"])
    for i, (a, e) in enumerate(zip(actual["body"], expected["body"])):
        assert a == e, f"body block {i} differs"


@pytest.mark.parametrize("name", ["CV_AdrianMatei.pdf", "CV_JohnDoe.pdf"])
def test_pdf_matches_golden(legacy_run, name):
    work, _ = legacy_run
    assert pdf_pages(work / "output" / "sample" / name) == pdf_pages(GOLDEN / name)


def test_legacy_path_never_imports_ai_packages(tmp_path):
    """generate.py and cvgen.render must import and run with the AI packages missing."""
    work = _workdir(tmp_path)
    blocker = (
        "import sys, runpy\n"
        f"BLOCKED = {AI_MODULES!r}\n"
        "class Block:\n"
        "    def find_spec(self, name, path=None, target=None):\n"
        "        if name.split('.')[0] in BLOCKED:\n"
        "            raise ImportError('blocked for test: ' + name)\n"
        "sys.meta_path.insert(0, Block())\n"
        f"sys.path.insert(0, {str(REPO)!r})\n"
        "import cvgen.render\n"
        f"sys.argv = [{str(REPO / 'generate.py')!r}]\n"
        f"runpy.run_path({str(REPO / 'generate.py')!r}, run_name='__main__')\n"
        "assert not any(m.split('.')[0] in BLOCKED for m in sys.modules)\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", blocker],
        cwd=work,
        env=_legacy_env(),
        capture_output=True,
        text=True,
        timeout=300,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert (work / "output" / "sample" / "CV_AdrianMatei.pdf").is_file()
